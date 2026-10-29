"""Fusionne la feuille du Copy Center avec les lignes Copy des listes fournisseur.

Une fiche par article physique : un article present des deux cotes devient UNE
ligne qui porte le nom, le code-barres et le prix du comptoir, et le libelle
fournisseur en description d'achat. Rien ne disparait : un article d'un seul
cote reste ou entre. Les papiers de copie sont regroupes par type, la marque
ne compte pas (decision de Mohamed, 2026-10-29).

Le rapprochement automatique ne suffit pas (0 nom identique, 0 code-barres
commun) : `review` ecrit un classeur ou une personne coche « meme article /
article different » ; `build` suit ces coches.

    uv run --no-project --with openpyxl python -I tools/fusion_copy_center.py review \\
        <feuille.xlsx> <Seed_Catalogue_v2.csv> <Product_Catalogue REVIEW.xlsx> <revue.xlsx>
    uv run ... tools/fusion_copy_center.py build \\
        <feuille.xlsx> <Seed_Catalogue_v2.csv> <Product_Catalogue REVIEW.xlsx> <revue.xlsx> <sortie.csv>
    uv run ... tools/fusion_copy_center.py --test

Ne touche pas a Odoo. La gouvernance MDM reste celle de l'import.
"""

import csv
import re
import sys
import unicodedata

BUREAUTIQUE = "Copy/Articles Bureautique"
FLEXY = "Copy/Flexy"
FIRST_NEW_REF = 1167  # le classeur du 2026-10-25 a attribue INV-000001..001166
SURE, DOUBT = 0.75, 0.4

# Les articles papier que consomment les copies : un par type, toutes marques
# confondues. Regles ancrees sur le debut du nom, sinon « Bloc Note A5 Luxe »
# deviendrait du papier de luxe.
PAPER_GROUPS = [
    ("Papier A4 80g", r"^rame .*\ba4\b|^papier noir/blanc$"),
    ("Papier Luxe", r"^papier (de )?luxe|^papier premium|prisma|ivo?i?ry"),
    ("Papier Bristol", r"^papier bristol"),
    ("Papier Autocollant", r"^(papier )?autocollant"),
    ("Papier Glacé", r"^papier glac"),
    ("Papier Photo", r"^papier photo"),
    ("Papier A5", r"^papier (extra blanc )?a5"),
]

# Les services de copie sont dans tools/seed/copy_services.csv, pas ici.
SERVICE = re.compile(r"^(impression|photocopie)\b", re.I)

COLUMNS = [
    "Reference",
    "Nom",
    "Ancien_Noms",
    "Categorie_Suggeree",
    "Type",
    "Code_Barres",
    "Prix_Vente",
    "Cout",
    "Description_Achat",
    "Categorie_POS",
    "Groupe_Papier",
    "Source",
    "Ligne",
]


# --------------------------------------------------------------- texte


def ascii_lower(text):
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text.lower()).strip()


def norm(text):
    return re.sub(r"[^a-z0-9]+", " ", ascii_lower(text)).strip()


def tokens(text):
    return {
        t for t in re.findall(r"[a-z]{3,}|\d+", ascii_lower(text)) if t not in {"pages", "avec", "pour", "des", "les"}
    }


def score(a, b):
    ta, tb = tokens(a), tokens(b)
    return len(ta & tb) / len(ta | tb) if ta and tb else 0.0


def paper_group(name):
    if SERVICE.match(str(name or "").strip()):
        return None
    low = ascii_lower(name)
    for group, pattern in PAPER_GROUPS:
        if re.search(pattern, low):
            return group
    return None


def barcode(value):
    value = str(value or "").strip()
    return "" if value in {"", "-"} else value


def money(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    return "" if number <= 0 else ("%g" % number)


# ------------------------------------------------------------- lecture


def read_sources(sheet_path, seed_path, catalogue_path):
    import openpyxl

    wb = openpyxl.load_workbook(sheet_path, data_only=True, read_only=True)

    def rows(tab):
        return list(wb[tab].iter_rows(values_only=True))[1:]

    counter = []
    for i, r in enumerate(rows("stock C.C."), start=2):
        if r[0] and str(r[0]).strip():
            counter.append(
                {
                    "id": "CC:%d" % i,
                    "name": str(r[0]).strip(),
                    "barcode": barcode(r[1]),
                    "price": money(r[2]),
                    "cost": money(r[3]),
                    "type": r[5],
                    "pos": (r[6] or "").strip(),
                }
            )
    shelf = []
    for i, r in enumerate(rows("Products"), start=2):
        if r[1] and str(r[1]).strip():
            shelf.append({"id": "PR:%d" % i, "name": str(r[1]).strip(), "barcode": barcode(r[2]), "pos": ""})
    old = []
    for i, r in enumerate(rows("Stock"), start=2):
        if r[1] and str(r[1]).strip():
            old.append(
                {
                    "id": "ST:%d" % i,
                    "name": str(r[1]).strip(),
                    "barcode": barcode(r[2]),
                    "price": money(r[3]),
                    "cost": money(r[4]),
                    "pos": (r[7] or "").strip(),
                }
            )
    with open(seed_path, encoding="utf-8-sig", newline="") as f:
        seed = [
            {"id": "SEED:%s:%s" % (r["Source"], r["Ligne"]), "name": r["Nom"].strip(), **r}
            for r in csv.DictReader(f)
            if r["Categorie_Suggeree"].startswith("Copy/")
        ]
    refs = {}
    cwb = openpyxl.load_workbook(catalogue_path, data_only=True, read_only=True)
    for r in list(cwb["Products"].iter_rows(values_only=True))[1:]:
        if r and r[0] and r[1]:
            refs[str(r[1]).strip()] = str(r[0]).strip()
    return counter, shelf, old, seed, refs


# --------------------------------------------------------- rapprochement


def best_pairs(left, right):
    """Chaque ligne de droite sert au plus une fois ; les meilleurs scores d'abord."""
    candidates = []
    for a in left:
        for b in right:
            if (a.get("barcode") and a["barcode"] == b.get("barcode")) or norm(a["name"]) == norm(b["name"]):
                s = 1.0
            else:
                s = score(a["name"], b["name"])
            if s >= DOUBT:
                candidates.append((s, a["id"], b["id"]))
    candidates.sort(key=lambda c: (-c[0], c[1], c[2]))
    used_a, used_b, pairs = set(), set(), {}
    for s, a, b in candidates:
        if a in used_a or b in used_b:
            continue
        used_a.add(a)
        used_b.add(b)
        pairs[a] = (b, s)
    return pairs


def analyse(counter, shelf, old, seed):
    """Tout ce que la revue montre et que build suit, sans les coches."""
    # Le comptage d'etagere et l'ancien export decrivent souvent le meme
    # article que le comptoir sous un autre nom (« stylo techno blue »,
    # « Cahier 120 Pages »). Sur rapprochement sur : ils completent la ligne du
    # comptoir (code-barres, cout). Sur rapprochement douteux : montres avec
    # leur jumeau probable, ecartes par defaut. Sans ressemblance : a importer.
    sheet = [c for c in counter if not SERVICE.match(c["name"]) and not re.search(r"\+\s*20", c["name"])]
    sheet_by_id = {c["id"]: c for c in sheet}
    others = []
    for extra_rows in (shelf, old):
        extra_rows = [r for r in extra_rows if not SERVICE.match(r["name"]) and not re.search(r"\+\s*20", r["name"])]
        twins = best_pairs(extra_rows, sheet)
        for r in extra_rows:
            hit = twins.get(r["id"])
            if hit and hit[1] >= SURE:
                target = sheet_by_id[hit[0]]
                target["barcode"] = target["barcode"] or r.get("barcode", "")
                target["cost"] = target.get("cost") or r.get("cost", "")
                continue
            twin = sheet_by_id[hit[0]]["name"] if hit else ""
            others.append(
                {
                    **r,
                    "price": r.get("price", ""),
                    "cost": r.get("cost", ""),
                    "twin": twin,
                    "twin_score": round(hit[1], 2) if hit else 0,
                    "default": "DROP" if hit else "IMPORT",
                }
            )
    stock_only = others

    groups = {}
    for row in sheet + stock_only + seed:
        group = paper_group(row["name"])
        if group:
            groups.setdefault(group, []).append(row)
    grouped = {row["id"] for members in groups.values() for row in members}

    left = [s for s in sheet if s["id"] not in grouped]
    seed_rest = [s for s in seed if s["id"] not in grouped]
    pairs = best_pairs(left, seed_rest)
    stock_left = [o for o in stock_only if o["id"] not in grouped]
    # Une ligne de l'ancien export peut aussi etre un article des listes.
    taken = {b for b, _s in pairs.values()}
    stock_pairs = best_pairs(stock_left, [s for s in seed_rest if s["id"] not in taken])

    seen, dup = {}, set()
    for row in left + stock_left:
        if row.get("barcode"):
            if row["barcode"] in seen:
                dup.add(row["barcode"])
            seen[row["barcode"]] = row
    return {
        "sheet": left,
        "seed": seed_rest,
        "stock_only": stock_left,
        "groups": groups,
        "pairs": pairs,
        "stock_pairs": stock_pairs,
        "dup_barcodes": dup,
    }


def default_decision(s):
    return "SAME" if s >= SURE else "?"


# -------------------------------------------------------------- revue


def write_review(path, data):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    seed_by_id = {s["id"]: s for s in data["seed"]}
    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "Read me"
    for line in [
        "Copy Center: merge review (counter sheet + supplier lists)",
        "",
        "Each physical article must exist once. Fill the Decision column, then send the file back.",
        "Pairs: SAME = one article (counter name, barcode and price kept, supplier name kept for purchases).",
        "       DIFF = two different articles. ? = not decided yet (treated as DIFF).",
        "Paper groups: IN = this paper is the copy paper of the group (brands do not matter). OUT = a separate article.",
        "Other tabs: rows of the Products and Stock tabs the counter list does not have.",
        "       IMPORT = a real article, create it. DROP = the same as 'Looks like', or no longer sold.",
        "To fix: problems to correct in the source sheet.",
    ]:
        readme.append([line])

    def tab(title, header, rows):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) for c in col)))
        return ws

    pairs = data["pairs"]
    rows = []
    for s in data["sheet"]:
        hit = pairs.get(s["id"])
        if hit:
            cand = seed_by_id[hit[0]]
            rows.append(
                [
                    s["id"],
                    s["name"],
                    s["barcode"],
                    s["price"],
                    hit[0],
                    cand["name"],
                    round(hit[1], 2),
                    default_decision(hit[1]),
                ]
            )
    for o in data["stock_only"]:
        hit = data["stock_pairs"].get(o["id"])
        if hit:
            cand = seed_by_id[hit[0]]
            rows.append(
                [
                    o["id"],
                    o["name"],
                    o["barcode"],
                    o["price"],
                    hit[0],
                    cand["name"],
                    round(hit[1], 2),
                    default_decision(hit[1]),
                ]
            )
    ws = tab(
        "Pairs", ["Sheet id", "Counter name", "Barcode", "Price", "Seed id", "Supplier name", "Score", "Decision"], rows
    )
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        if row[7].value == "?":
            row[7].fill = yellow

    tab(
        "Paper groups",
        ["Group", "Member id", "Member name", "Decision"],
        [[g, m["id"], m["name"], "IN"] for g, members in data["groups"].items() for m in members],
    )
    paired_sheet = set(pairs)
    tab(
        "Sheet only",
        ["Sheet id", "Counter name", "Barcode", "Price", "POS category"],
        [
            [s["id"], s["name"], s["barcode"], s["price"], s["pos"]]
            for s in data["sheet"]
            if s["id"] not in paired_sheet
        ],
    )
    used_seed = {b for b, _s in pairs.values()} | {b for b, _s in data["stock_pairs"].values()}
    tab(
        "Seed only",
        ["Seed id", "Supplier name"],
        [[s["id"], s["name"]] for s in data["seed"] if s["id"] not in used_seed],
    )
    tab(
        "Other tabs",
        ["Id", "Name (Products / Stock tab)", "Price", "Cost", "Looks like (counter)", "Score", "Decision"],
        [
            [o["id"], o["name"], o["price"], o["cost"], o["twin"], o["twin_score"], o["default"]]
            for o in data["stock_only"]
            if o["id"] not in data["stock_pairs"]
        ],
    )
    fixes = []
    for bc in sorted(data["dup_barcodes"]):
        names = [r["name"] for r in data["sheet"] + data["stock_only"] if r.get("barcode") == bc]
        fixes.append([bc, " / ".join(names), "Same barcode on several articles: imported WITHOUT barcode until fixed."])
    for s in data["sheet"]:
        if not s["price"]:
            fixes.append([s["id"], s["name"], "No sale price: imported but not sellable until priced."])
    tab("To fix", ["Where", "Article", "Problem"], fixes)
    wb.save(path)


def read_decisions(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    pairs, groups, stock = {}, {}, {}
    for r in list(wb["Pairs"].iter_rows(values_only=True))[1:]:
        if r and r[0]:
            pairs[r[0]] = (r[4], str(r[7] or "?").strip().upper())
    for r in list(wb["Paper groups"].iter_rows(values_only=True))[1:]:
        if r and r[1]:
            groups[r[1]] = str(r[3] or "IN").strip().upper()
    for r in list(wb["Other tabs"].iter_rows(values_only=True))[1:]:
        if r and r[0]:
            stock[r[0]] = str(r[6] or "IMPORT").strip().upper()
    return pairs, groups, stock


# --------------------------------------------------------------- build


def build_rows(data, refs, decisions):
    pair_dec, group_dec, stock_dec = decisions
    seed_by_id = {s["id"]: s for s in data["seed"]}
    out, used_seed, retired = [], set(), []

    def row_from(src, **extra):
        categ = FLEXY if src.get("pos") == "Flexy" else BUREAUTIQUE
        return {
            "Reference": "",
            "Nom": src["name"],
            "Ancien_Noms": "",
            "Categorie_Suggeree": categ,
            "Type": "Service" if categ == FLEXY else "Storable",
            "Code_Barres": "" if src.get("barcode") in data["dup_barcodes"] else src.get("barcode", ""),
            "Prix_Vente": src.get("price", ""),
            "Cout": src.get("cost", ""),
            "Description_Achat": "",
            "Categorie_POS": src.get("pos") or "Articles Bureautique",
            "Groupe_Papier": "",
            "Source": "copy_center_fusion",
            "Ligne": src["id"],
            **extra,
        }

    def merge_seed(row, seeds):
        codes = sorted(refs[s["name"]] for s in seeds if s["name"] in refs)
        if codes:
            row["Reference"] = codes[0]
            retired.extend(codes[1:])
        row["Ancien_Noms"] = " | ".join(s["name"] for s in seeds)
        row["Description_Achat"] = " | ".join(s["name"] for s in seeds)
        used_seed.update(s["id"] for s in seeds)
        return row

    def pair_of(row_id, auto):
        if row_id in pair_dec:
            seed_id, decision = pair_dec[row_id]
            return seed_by_id.get(seed_id) if decision == "SAME" else None
        hit = auto.get(row_id)
        return seed_by_id[hit[0]] if hit and default_decision(hit[1]) == "SAME" else None

    for s in data["sheet"]:
        twin = pair_of(s["id"], data["pairs"])
        out.append(merge_seed(row_from(s), [twin]) if twin else row_from(s))

    for o in data["stock_only"]:
        twin = pair_of(o["id"], data["stock_pairs"])
        if twin:
            out.append(merge_seed(row_from(o), [twin]))
        elif stock_dec.get(o["id"], o.get("default", "IMPORT")) != "DROP":
            out.append(row_from(o))

    for group, members in data["groups"].items():
        inside = [m for m in members if group_dec.get(m["id"], "IN") != "OUT"]
        outside = [m for m in members if m not in inside]
        for m in outside:
            if m["id"].startswith("SEED:"):
                continue  # repris plus bas avec les lignes fournisseur seules
            out.append(row_from(m))
        if not inside:
            continue
        seeds = [m for m in inside if m["id"].startswith("SEED:")]
        priced = [m.get("price") for m in inside if m.get("price")]
        costed = [m.get("cost") for m in inside if m.get("cost")]
        row = row_from(
            {"id": "GROUP:%s" % group, "name": group, "pos": "Articles Bureautique"},
            Groupe_Papier=group,
            Prix_Vente=max(priced, key=priced.count) if priced else "",
            Cout=max(costed, key=costed.count) if costed else "",
        )
        merge_seed(row, seeds)
        row["Description_Achat"] = " | ".join(m["name"] for m in inside)
        # Les codes-barres des marques ne designent plus un article : la rame
        # vendue telle quelle est un service de copy_services.csv.
        row["Code_Barres"] = ""
        out.append(row)
        # Une ligne fournisseur sortie du groupe redevient un article a part.
        data["seed"].extend(m for m in outside if m["id"].startswith("SEED:") and m not in data["seed"])

    for s in data["seed"]:
        if s["id"] in used_seed:
            continue
        out.append(
            {
                **row_from(
                    {"id": s["id"], "name": s["name"]}, Categorie_Suggeree=s["Categorie_Suggeree"], Type=s["Type"]
                ),
                "Reference": refs.get(s["name"], ""),
                "Ancien_Noms": s["name"],
                "Source": s["Source"],
                "Ligne": s["Ligne"],
                "Categorie_POS": "",
            }
        )
        used_seed.add(s["id"])

    # Meme nom = meme article, quelle que soit la feuille d'ou il vient (l'ancien
    # export repete « Note Book HTC »). La ligne qui porte deja un code du
    # classeur l'emporte ; l'autre lui donne ce qui lui manque.
    merged, conflicts = {}, []
    for row in sorted(out, key=lambda r: (not re.fullmatch(r"INV-\d{6}", r["Reference"] or ""), r["Ligne"])):
        key = norm(row["Nom"])
        keep = merged.get(key)
        if not keep:
            merged[key] = row
            continue
        if row["Code_Barres"] and keep["Code_Barres"] and row["Code_Barres"] != keep["Code_Barres"]:
            conflicts.append("%s: %s / %s" % (keep["Nom"], keep["Code_Barres"], row["Code_Barres"]))
        for field in ("Code_Barres", "Prix_Vente", "Cout", "Description_Achat", "Ancien_Noms", "Reference"):
            keep[field] = keep[field] or row[field]
        if row["Reference"] and row["Reference"] != keep["Reference"] and re.fullmatch(r"INV-\d{6}", row["Reference"]):
            retired.append(row["Reference"])
    out = list(merged.values())
    for line in conflicts:
        print("Same name, two barcodes (first kept): %s" % line)

    # Jamais de reference semantique, et les nouvelles a la suite du classeur.
    for row in out:
        if not re.fullmatch(r"INV-\d{6}", row["Reference"] or ""):
            row["Reference"] = ""
    nxt = FIRST_NEW_REF
    for row in sorted((r for r in out if not r["Reference"]), key=lambda r: (r["Categorie_Suggeree"], norm(r["Nom"]))):
        row["Reference"] = "INV-%06d" % nxt
        nxt += 1
    out.sort(key=lambda r: (r["Categorie_Suggeree"], norm(r["Nom"])))
    return out, sorted(set(retired))


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


# ---------------------------------------------------------------- tests


def _autotest():
    assert norm("Cahier 120 Pages  El-Hilal") == "cahier 120 pages el hilal"
    assert paper_group("RAME EXTRA A4 80G 500F TARGET") == "Papier A4 80g"
    assert paper_group("RAME DE PAPIER A4 / 500 FEUILLES – RAMOS ") == "Papier A4 80g"
    assert paper_group("Papier Noir/Blanc") == "Papier A4 80g"
    assert paper_group("RAME EXTRA A3 80G RAMOS") is None
    assert paper_group("Impression Papier Autocollant") is None
    assert paper_group("Bloc Note A5 Luxe") is None
    assert paper_group("PAPIER DE LUXE PRISMA A4 130 G 9517 TECHNO") == "Papier Luxe"
    assert paper_group("Papier Glacé") == "Papier Glacé"

    counter = [
        {
            "id": "CC:2",
            "name": "Stylo Ultron Noir",
            "barcode": "8904348051941",
            "price": "50",
            "cost": "",
            "type": None,
            "pos": "Articles Bureautique",
        },
        {
            "id": "CC:3",
            "name": "Cartouche Golden Rouge",
            "barcode": "6975345543281",
            "price": "25",
            "cost": "",
            "type": None,
            "pos": "",
        },
        {
            "id": "CC:4",
            "name": "Cartouche Golden Vert",
            "barcode": "6975345543281",
            "price": "25",
            "cost": "",
            "type": None,
            "pos": "",
        },
        {
            "id": "CC:5",
            "name": "Impression Noir/Blanc +20 Pages",
            "barcode": "",
            "price": "8",
            "cost": "",
            "type": "Service",
            "pos": "",
        },
        {
            "id": "CC:6",
            "name": "RAME DE PAPIER A4 / 500 FEUILLES – RAMOS",
            "barcode": "6131145000014",
            "price": "650",
            "cost": "",
            "type": None,
            "pos": "",
        },
    ]
    shelf = [{"id": "PR:2", "name": "ultron noir", "barcode": "8904348051941", "pos": ""}]
    old = [
        {"id": "ST:2", "name": "Agrafes", "barcode": "", "price": "50", "cost": "", "pos": ""},
        {"id": "ST:3", "name": "Papier Noir/Blanc", "barcode": "", "price": "2", "cost": "1.2", "pos": ""},
    ]
    seed = [
        {
            "id": "SEED:l:4",
            "name": "STYLO ULTRON X2 NOIR",
            "Categorie_Suggeree": BUREAUTIQUE,
            "Type": "Storable",
            "Source": "l",
            "Ligne": "4",
        },
        {
            "id": "SEED:l:5",
            "name": "STYLO ULTRON X2 BLEU",
            "Categorie_Suggeree": BUREAUTIQUE,
            "Type": "Storable",
            "Source": "l",
            "Ligne": "5",
        },
        {
            "id": "SEED:l:9",
            "name": "RAME EXTRA A4 80G 500F TARGET",
            "Categorie_Suggeree": BUREAUTIQUE,
            "Type": "Storable",
            "Source": "l",
            "Ligne": "9",
        },
        {
            "id": "SEED:l:10",
            "name": "RAME EXTRA A4 80G GOLDEN / RAMOS ROUGE",
            "Categorie_Suggeree": BUREAUTIQUE,
            "Type": "Storable",
            "Source": "l",
            "Ligne": "10",
        },
    ]
    refs = {
        "STYLO ULTRON X2 NOIR": "INV-000010",
        "RAME EXTRA A4 80G 500F TARGET": "INV-000030",
        "RAME EXTRA A4 80G GOLDEN / RAMOS ROUGE": "INV-000020",
        "STYLO ULTRON X2 BLEU": "COP-ABU-AGRF",
    }
    data = analyse(counter, shelf, old, seed)
    assert set(data["groups"]) == {"Papier A4 80g"}
    assert len(data["groups"]["Papier A4 80g"]) == 4  # rame du comptoir, ancien export, 2 listes
    assert "6975345543281" in data["dup_barcodes"]
    pairs = data["pairs"]
    assert pairs["CC:2"][0] == "SEED:l:4"  # le meme Ultron, pas le bleu
    assert len({b for b, _s in pairs.values()}) == len(pairs), "une ligne fournisseur ne sert qu'une fois"
    rows, retired = build_rows(data, refs, ({"CC:2": ("SEED:l:4", "SAME")}, {}, {}))
    by_name = {r["Nom"]: r for r in rows}
    assert by_name["Stylo Ultron Noir"]["Reference"] == "INV-000010"
    assert by_name["Stylo Ultron Noir"]["Description_Achat"] == "STYLO ULTRON X2 NOIR"
    assert by_name["Papier A4 80g"]["Reference"] == "INV-000020" and retired == ["INV-000030"]
    assert by_name["Papier A4 80g"]["Code_Barres"] == ""
    assert by_name["Cartouche Golden Rouge"]["Code_Barres"] == "" == by_name["Cartouche Golden Vert"]["Code_Barres"]
    assert not any("+20" in r["Nom"] for r in rows)
    assert not any(r["Reference"].startswith("COP-") for r in rows)
    assert by_name["STYLO ULTRON X2 BLEU"]["Reference"].startswith("INV-0011")  # COP- ecarte, nouveau code
    assert min(r["Reference"] for r in rows if r["Reference"] >= "INV-001167") == "INV-001167"
    assert "Agrafes" in by_name  # ancien export sans ressemblance : importe par defaut
    data2 = analyse(
        [dict(counter[0], name="Cahier 120 Pages El-Hilal", id="CC:9", barcode="")],
        [],
        [{"id": "ST:9", "name": "Cahier 120 Pages", "barcode": "", "price": "140", "cost": "", "pos": ""}],
        [],
    )
    assert data2["stock_only"][0]["default"] == "DROP" and data2["stock_only"][0]["twin"] == "Cahier 120 Pages El-Hilal"
    rows2, _r = build_rows(data2, {}, ({}, {}, {}))
    assert [r["Nom"] for r in rows2] == ["Cahier 120 Pages El-Hilal"]  # pas de doublon par defaut
    data3 = analyse(
        [],
        [],
        [
            {"id": "ST:7", "name": "Note Book HTC", "barcode": "", "price": "", "cost": "380", "pos": ""},
            {"id": "ST:8", "name": "Note Book HTC", "barcode": "", "price": "", "cost": "380", "pos": ""},
        ],
        [],
    )
    rows3, _r = build_rows(data3, {}, ({}, {}, {}))
    assert len(rows3) == 1  # meme nom, un seul article
    print("fusion_copy_center : autotest OK")


def main(argv):
    if argv[1:] == ["--test"]:
        _autotest()
        return
    command, *paths = argv[1:]
    sheet, seed, catalogue = paths[:3]
    *sources, refs = read_sources(sheet, seed, catalogue)
    data = analyse(*sources)
    if command == "review":
        write_review(paths[3], data)
        sure = sum(1 for _b, s in data["pairs"].values() if s >= SURE)
        print("Pairs: %d sure, %d to decide" % (sure, len(data["pairs"]) - sure))
        print("Paper groups: %s" % {g: len(m) for g, m in data["groups"].items()})
        print(
            "Sheet rows: %d, seed rows: %d, stock tab only: %d"
            % (len(data["sheet"]), len(data["seed"]), len(data["stock_only"]))
        )
        print("Duplicate barcodes: %s" % sorted(data["dup_barcodes"]))
    elif command == "build":
        rows, retired = build_rows(data, refs, read_decisions(paths[3]))
        write_csv(paths[4], rows)
        undecided = sum(1 for _s, d in read_decisions(paths[3])[0].values() if d == "?")
        print("Rows written: %d (%d undecided pairs kept apart)" % (len(rows), undecided))
        print("Retired references: %s" % (", ".join(retired) or "none"))
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv)
